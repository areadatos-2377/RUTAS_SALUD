from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from datetime import date

import openpyxl
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from .management.commands.cargar_clues import Command, fecha_programacion_celda
from .models import UnidadMedica


class CargarCluesContactosTests(TestCase):
	def test_importa_contactos_desde_columnas_del_catalogo(self):
		with TemporaryDirectory() as directorio:
			clues_path = Path(directorio) / "CLUES_IMB.xlsx"
			entidades_path = Path(directorio) / "entidades.xlsx"

			wb_clues = openpyxl.Workbook()
			ws_clues = wb_clues.active
			ws_clues.title = "BD_IMB"
			ws_clues.append(
				[
					"NIVEL ATENCION",
					"ESTATUS DE OPERACION",
					"ENTIDAD",
					"CLUES",
					"NOMBRE DE TIPOLOGIA",
					"NOMBRE DE LA UNIDAD",
					"MUNICIPIO",
					"¿QUIÉN RECIBE EN UNIDAD?",
					"TELÉFONO",
					"CORREO",
					"RUTA",
					"FECHA PROG",
				]
			)
			ws_clues.append(
				[
					"PRIMER NIVEL",
					"EN OPERACION",
					"COLIMA",
					"CMIMB000001",
					"CENTRO DE SALUD",
					"Unidad de prueba",
					"Colima",
					"Responsable de farmacia",
					"312 123 4567 EXT 8",
					"farmacia@ejemplo.test - almacen@ejemplo.test",
					4,
					"Thursday, August 20, 2026",
				]
			)
			wb_clues.save(clues_path)

			wb_entidades = openpyxl.Workbook()
			ws_entidades = wb_entidades.active
			ws_entidades.title = "Hoja1"
			ws_entidades.append([])
			ws_entidades.append([])
			ws_entidades.append([])
			ws_entidades.append(["COLIMA", "Coordinador"])
			wb_entidades.save(entidades_path)

			with (
				patch(
					"catalogos.management.commands.cargar_clues.CLUES_XLSX",
					clues_path,
				),
				patch(
					"catalogos.management.commands.cargar_clues.ENTIDADES_XLSX",
					entidades_path,
				),
			):
				call_command("cargar_clues", verbosity=0)

		unidad = UnidadMedica.objects.get(pk="CMIMB000001")
		self.assertEqual(unidad.quien_recibe, "Responsable de farmacia")
		self.assertEqual(unidad.telefono, "312 123 4567 EXT 8")
		self.assertEqual(
			unidad.correo,
			"farmacia@ejemplo.test - almacen@ejemplo.test",
		)
		self.assertEqual(unidad.ruta_programacion, "4")
		self.assertEqual(unidad.fecha_programacion_referencia, date(2026, 8, 20))

	def test_falla_con_mensaje_claro_si_faltan_columnas_de_programacion(self):
		with TemporaryDirectory() as directorio:
			clues_path = Path(directorio) / "CLUES_IMB.xlsx"
			wb = openpyxl.Workbook()
			ws = wb.active
			ws.title = "BD_IMB"
			ws.append(["CLUES"])
			wb.save(clues_path)

			with (
				patch(
					"catalogos.management.commands.cargar_clues.CLUES_XLSX",
					clues_path,
				),
				patch.object(Command, "_leer_coordinadores", return_value=({}, [])),
				self.assertRaisesMessage(CommandError, "FECHA PROG"),
			):
				Command()._cargar(False, 0, False)

	def test_fecha_no_reconocida_se_conserva_vacia(self):
		self.assertIsNone(fecha_programacion_celda("SIN FECHA DE PROGRAMACIÓN"))
