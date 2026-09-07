"""Genera la presentacion de PowerPoint de evidencia de picking & packing, a
partir de la plantilla real (picking_packing/plantillas/Formato_picking_packing.pptx,
copia de docs/Formato_picking_packing.pptx, que a su vez viene de
legacy/Picking Packing.pptx). Mismo mecanismo que entregas/presentacion.py
(ver ese modulo y entregas/pptx_utils.py para el detalle de por que se hace
asi con python-pptx) -- la diferencia de formato es que aqui cada
diapositiva de contenido trae 4 ENTIDADES (no una), cada una con un par de
slots FOTO + VIDEO (no 8 fotos de una sola entidad).

El video se inserta embebido y reproducible (Shapes.add_movie) sin poster
frame real -- python-pptx usa su imagen default de "media loudspeaker" como
cartel. Generar un fotograma real requeriria ffmpeg en el servidor, que se
decidio no agregar (ver el plan de esta funcion).
"""

import io
import logging
from pathlib import Path

from pptx import Presentation

from entregas import storage
from entregas.pptx_utils import duplicar_diapositiva, eliminar_diapositiva, forma_por_id, texto_forma
from entregas.regiones import orden_regiones, region_de

logger = logging.getLogger(__name__)

RUTA_PLANTILLA = Path(__file__).resolve().parent / "plantillas" / "Formato_picking_packing.pptx"

_MIME_POR_EXTENSION = {".mp4": "video/mp4", ".mov": "video/quicktime"}

# IDs de forma tomados de la plantilla real (ver el plan de esta funcion) --
# cada una de las 4 celdas de la diapositiva de contenido trae un titulo
# (nombre de entidad) y un par de rectangulos FOTO/VIDEO, que se sustituyen
# por una Picture/Movie real en la misma posicion (EMU) y se borran si no
# hay entidad o archivo para ese slot.
_CELDAS = [
    {"caption": 17, "foto": 2, "video": 4},
    {"caption": 31, "foto": 7, "video": 8},
    {"caption": 35, "foto": 10, "video": 11},
    {"caption": 38, "foto": 12, "video": 13},
]

_ID_DIA_CONTENIDO = 9
_ID_TITULO_REGION = 2
_ID_DIA_PORTADA = 16
_ID_TITULO_PORTADA = 15


def _texto_titulo_portada(forma, texto_nuevo):
    """El titulo de la portada trae 3 runs en un solo parrafo: 'Evidencia
    fotografica y de video:' / '«Picking y packing» para <categoria> de ' /
    'atencion medica' -- solo se cambia el run de en medio (la categoria),
    igual que _texto_subtitulo_portada en entregas/presentacion.py."""
    forma.text_frame.paragraphs[0].runs[1].text = texto_nuevo


def _insertar_foto(diapositiva, posicion, evidencia):
    bytes_imagen = storage.descargar_evidencia(evidencia.ruta_almacen)
    left, top, ancho, alto = posicion
    diapositiva.shapes.add_picture(io.BytesIO(bytes_imagen), left, top, ancho, alto)


def _insertar_video(diapositiva, posicion, evidencia):
    bytes_video = storage.descargar_evidencia(evidencia.ruta_almacen)
    left, top, ancho, alto = posicion
    ext = storage.extension(evidencia.nombre_original)
    mime_type = _MIME_POR_EXTENSION.get(ext, "video/unknown")
    diapositiva.shapes.add_movie(io.BytesIO(bytes_video), left, top, ancho, alto, mime_type=mime_type)


def _reemplazar_rectangulo(diapositiva, forma_slot, evidencia, insertar, entidad_nombre, etiqueta):
    """El rectangulo 'FOTO'/'VIDEO' de la plantilla sigue presente (a
    diferencia de Formato_rutas.pptx, donde ya se habian quitado) -- hay que
    borrarlo siempre, tanto si se inserta la Picture/Movie real en su lugar
    como si el slot queda vacio, o el rectangulo de muestra se quedaria
    encima/al lado en la diapositiva final."""
    posicion = (forma_slot.left, forma_slot.top, forma_slot.width, forma_slot.height)
    forma_slot._element.getparent().remove(forma_slot._element)
    if evidencia is None:
        return
    try:
        insertar(diapositiva, posicion, evidencia)
    except Exception:
        logger.warning(
            "No se pudo insertar el/la %s de picking&packing id=%s (entidad %s)",
            etiqueta, evidencia.id, entidad_nombre, exc_info=True,
        )


def _llenar_celda(diapositiva, celda, item):
    forma_caption = forma_por_id(diapositiva, celda["caption"])
    forma_foto = forma_por_id(diapositiva, celda["foto"])
    forma_video = forma_por_id(diapositiva, celda["video"])

    if item is None:
        for forma in (forma_caption, forma_foto, forma_video):
            forma._element.getparent().remove(forma._element)
        return

    texto_forma(forma_caption, item["entidad"].nombre.title())
    _reemplazar_rectangulo(diapositiva, forma_foto, item["foto"], _insertar_foto, item["entidad"].nombre, "foto")
    _reemplazar_rectangulo(diapositiva, forma_video, item["video"], _insertar_video, item["entidad"].nombre, "video")


def construir_presentacion(dia_texto, categoria_texto, entidades):
    """entidades: lista de {"entidad": Entidad, "foto": Evidencia|None,
    "video": Evidencia|None} -- ya resuelta y validada desde la base (ver
    views.py), una entrada por cada entidad con evidencia ese dia.
    categoria_texto es jornada.get_categoria_display() ("Primer nivel" o
    "Segundo y tercer nivel"). Regresa un BytesIO listo para mandar como
    respuesta binaria."""
    prs = Presentation(str(RUTA_PLANTILLA))
    diapositiva_portada = prs.slides[0]
    diapositiva_region_base = prs.slides[1]
    diapositiva_contenido_base = prs.slides[2]

    dia_etiqueta = f"Día 1: {dia_texto}"
    texto_forma(forma_por_id(diapositiva_portada, _ID_DIA_PORTADA), dia_etiqueta)
    _texto_titulo_portada(
        forma_por_id(diapositiva_portada, _ID_TITULO_PORTADA),
        f"«Picking y packing» para {categoria_texto.lower()} de ",
    )

    # Region (orden fijo) -> entidad (orden alfabetico) -> item de esa entidad.
    por_region = {}
    for item in entidades:
        region = region_de(item["entidad"].nombre)
        por_region.setdefault(region, {})[item["entidad"].nombre] = item

    # diapositiva_region_base y diapositiva_contenido_base son PLANTILLAS --
    # nunca se editan/mutan directamente (mismo motivo que en
    # entregas/presentacion.py: llenar una copia con menos de 4 entidades
    # borra celdas; si le pasara a la base, la siguiente copia saldria
    # incompleta). Al final se borran las 2 plantillas.
    for region in orden_regiones():
        entidades_region = por_region.get(region)
        if not entidades_region:
            continue

        diapositiva_region = duplicar_diapositiva(prs, diapositiva_region_base)
        texto_forma(forma_por_id(diapositiva_region, _ID_TITULO_REGION), region)

        items_ordenados = [entidades_region[nombre] for nombre in sorted(entidades_region)]
        for inicio in range(0, len(items_ordenados), len(_CELDAS)):
            grupo = items_ordenados[inicio: inicio + len(_CELDAS)]
            diapositiva_contenido = duplicar_diapositiva(prs, diapositiva_contenido_base)
            texto_forma(forma_por_id(diapositiva_contenido, _ID_DIA_CONTENIDO), dia_etiqueta)
            for i, celda in enumerate(_CELDAS):
                item = grupo[i] if i < len(grupo) else None
                _llenar_celda(diapositiva_contenido, celda, item)

    eliminar_diapositiva(prs, diapositiva_region_base)
    eliminar_diapositiva(prs, diapositiva_contenido_base)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
