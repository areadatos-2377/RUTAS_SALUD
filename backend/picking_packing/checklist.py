"""Genera el checklist de picking & packing en Excel (una pestaña por dia
del periodo de la jornada, igual que legacy/Checklist_Picking_1er
Nivel_5a entrega.xlsx). A diferencia de la presentacion PPTX, aqui NO se
parte de la plantilla real: el archivo original usa una casilla de
verificacion nativa de Excel (funcion 2024, ligada a xl/featurePropertyBag)
que openpyxl no puede preservar al guardar -- se probo de forma directa
(round-trip: abrir con openpyxl, guardar, esa pieza desaparece). En vez de
manipular el XML del .xlsx a mano (fragil, sin forma de verificar
visualmente en este entorno que Excel lo siga abriendo bien), se construye
con la API normal de openpyxl y un check de texto (✓) en vez de una
casilla clicable -- decision tomada con el usuario, ver el plan.
"""

import io
from datetime import timedelta
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as ImagenXLSX
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from entregas.fechas import formato_fecha_es
from entregas.regiones import orden_regiones, region_de

# Copia de legacy/logo rutas.png -- vive dentro de backend/ (no en docs/ ni
# legacy/) porque el servicio de Railway se construye usando backend/ como
# raiz aislada (ver backend/README.md, seccion del catalogo de CLUES);
# cualquier archivo fuera de esa carpeta no existe en el contenedor.
_RUTA_LOGO = Path(__file__).resolve().parent / "plantillas" / "logo_rutas.png"

_VERDE_BANNER = "FF338571"
_GRIS_HEADER = "FFBACEC6"
_VERDE_CHECK = "FF1E7A46"
_GRIS_BORDE = "FFCBD5D1"

_BORDE = Border(*(Side(style="thin", color=_GRIS_BORDE) for _ in range(4)))
_COLUMNAS = ["Región", "Entidad", "Foto", "Video", "Observaciones"]
_ANCHOS = [16, 24, 10, 10, 26]


def _rango_dias(inicio, fin):
    dias = []
    actual = inicio
    while actual <= fin:
        dias.append(actual)
        actual += timedelta(days=1)
    return dias


def _llenar_hoja(ws, titulo_l1, titulo_l2, entidades_por_region, marcas_por_entidad):
    for i, ancho in enumerate(_ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.merge_cells("A1:E1")
    banner = ws["A1"]
    banner.value = f"{titulo_l1}\n{titulo_l2}"
    banner.font = Font(name="Calibri", size=12, color="FFFFFFFF")
    banner.fill = PatternFill(fill_type="solid", fgColor=_VERDE_BANNER)
    banner.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42

    # Un objeto Image no se puede reutilizar entre hojas (openpyxl lo liga a
    # la hoja donde se agrega) -- se crea uno nuevo por cada llamada. Se
    # ancla en la esquina superior derecha del banner, mismo lugar que en
    # legacy/Checklist_Picking_1er Nivel_5a entrega.xlsx.
    logo = ImagenXLSX(str(_RUTA_LOGO))
    alto_original, ancho_original = logo.height, logo.width
    logo.height = 40
    logo.width = round(logo.height * ancho_original / alto_original)
    ws.add_image(logo, "E1")

    for i, encabezado in enumerate(_COLUMNAS, start=1):
        celda = ws.cell(row=2, column=i, value=encabezado)
        celda.font = Font(name="Calibri", size=11, bold=True)
        celda.fill = PatternFill(fill_type="solid", fgColor=_GRIS_HEADER)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = _BORDE

    fila = 3
    for region in orden_regiones():
        entidades = entidades_por_region.get(region)
        if not entidades:
            continue

        fila_inicio_region = fila
        for entidad in entidades:
            marcas = marcas_por_entidad.get(entidad.id, {"foto": False, "video": False})

            ws.cell(row=fila, column=1).border = _BORDE
            c_entidad = ws.cell(row=fila, column=2, value=entidad.nombre.title())
            c_entidad.alignment = Alignment(horizontal="left", vertical="center")
            c_entidad.border = _BORDE

            for columna, tipo in ((3, "foto"), (4, "video")):
                c = ws.cell(row=fila, column=columna, value="✓" if marcas.get(tipo) else "")
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = Font(name="Calibri", size=12, bold=True, color=_VERDE_CHECK)
                c.border = _BORDE

            c_obs = ws.cell(row=fila, column=5, value="")
            c_obs.border = _BORDE

            fila += 1

        c_region = ws.cell(row=fila_inicio_region, column=1, value=region)
        c_region.font = Font(name="Calibri", size=11, bold=True)
        c_region.alignment = Alignment(horizontal="center", vertical="center")
        if fila - 1 > fila_inicio_region:
            ws.merge_cells(start_row=fila_inicio_region, start_column=1, end_row=fila - 1, end_column=1)

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


def construir_checklist(jornada, entidades, evidencia_por_dia):
    """entidades: lista de catalogos.Entidad que participan en la jornada.
    evidencia_por_dia: {date: {entidad_id: {"foto": bool, "video": bool}}}
    (ya resuelto desde la base, ver views.py). Regresa un BytesIO listo
    para mandar como respuesta binaria."""
    entidades_por_region = {}
    for entidad in entidades:
        entidades_por_region.setdefault(region_de(entidad.nombre), []).append(entidad)
    for lista in entidades_por_region.values():
        lista.sort(key=lambda e: e.nombre)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    dias = _rango_dias(jornada.fecha_inicio, jornada.fecha_fin)
    for indice, dia in enumerate(dias, start=1):
        ws = wb.create_sheet(f"Día {indice}")
        marcas = evidencia_por_dia.get(dia, {})
        _llenar_hoja(ws, jornada.nombre, f"Picking y packing / {formato_fecha_es(dia)}", entidades_por_region, marcas)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
