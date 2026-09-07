"""Helpers de python-pptx compartidos entre entregas/presentacion.py y
picking_packing/presentacion.py -- ninguno es especifico del formato de
Distribucion, son mecanica pura de python-pptx (duplicar/borrar
diapositivas, cambiar texto preservando formato). Separado para no
duplicar esto entre los dos generadores de PPTX."""

import copy

from pptx.opc.constants import RELATIONSHIP_TYPE as RT

_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def texto_forma(forma, texto_nuevo):
    """Cambia el texto de una forma preservando el formato (fuente, tamano,
    color) del primer run -- reemplazar text_frame.text de un jalon lo
    resetea a la fuente por default."""
    parrafo = forma.text_frame.paragraphs[0]
    if not parrafo.runs:
        parrafo.add_run()
    parrafo.runs[0].text = texto_nuevo
    for run_extra in parrafo.runs[1:]:
        run_extra.text = ""


def forma_por_id(diapositiva, shape_id):
    for forma in diapositiva.shapes:
        if forma.shape_id == shape_id:
            return forma
    raise ValueError(f"No se encontro la forma id={shape_id} en la diapositiva -- "
                      "revisar si la plantilla cambio de estructura.")


def _remapear_imagenes(elemento, parte_origen, parte_nueva):
    for nodo in elemento.iter():
        for atributo in ("embed", "link"):
            rid_viejo = nodo.get(f"{_NS_REL}{atributo}")
            if not rid_viejo:
                continue
            parte_imagen = parte_origen.related_part(rid_viejo)
            rid_nuevo = parte_nueva.relate_to(parte_imagen, RT.IMAGE)
            nodo.set(f"{_NS_REL}{atributo}", rid_nuevo)


def eliminar_diapositiva(prs, diapositiva):
    """python-pptx no trae 'borrar diapositiva' -- hay que quitar su entrada
    de sldIdLst y soltar la relacion en el part de la presentacion, si no
    el archivo queda con una referencia rota."""
    for id_slide in list(prs.slides._sldIdLst):
        if prs.part.related_part(id_slide.rId) is diapositiva.part:
            prs.part.drop_rel(id_slide.rId)
            prs.slides._sldIdLst.remove(id_slide)
            return


def duplicar_diapositiva(prs, diapositiva_origen):
    nueva = prs.slides.add_slide(diapositiva_origen.slide_layout)
    # add_slide() trae los placeholders vacios del layout -- no sirven,
    # se quitan antes de copiar encima las formas reales.
    for forma in list(nueva.shapes):
        forma._element.getparent().remove(forma._element)
    for forma_origen in diapositiva_origen.shapes:
        copia = copy.deepcopy(forma_origen._element)
        _remapear_imagenes(copia, diapositiva_origen.part, nueva.part)
        nueva.shapes._spTree.append(copia)
    return nueva
